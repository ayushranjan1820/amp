import os
import re
import math
import base64
import uuid
import time
import requests
import io
from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime
from collections import Counter


def _extract_images_from_pdf(content_bytes: bytes) -> List[Dict[str, Any]]:
    extracted = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content_bytes))
        for page_num, page in enumerate(reader.pages, 1):
            try:
                for img_idx, image in enumerate(page.images):
                    img_bytes = image.data
                    if len(img_bytes) < 200:
                        continue
                    img_b64 = base64.b64encode(img_bytes).decode("utf-8")
                    name = getattr(image, "name", f"image_{page_num}_{img_idx}")
                    ext = name.rsplit(".", 1)[-1].lower() if "." in name else "png"
                    mime = {
                        "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                        "gif": "image/gif", "bmp": "image/bmp", "tiff": "image/tiff",
                        "webp": "image/webp",
                    }.get(ext, "image/png")
                    extracted.append({
                        "page": page_num,
                        "index": img_idx,
                        "name": name,
                        "image_b64": img_b64,
                        "mime_type": mime,
                        "size_bytes": len(img_bytes),
                        "type": "pdf_image",
                    })
            except Exception as e:
                print(f"[KB Ingestion] Error extracting images from PDF page {page_num}: {e}")
    except Exception as e:
        print(f"[KB Ingestion] pypdf image extraction failed: {e}")
    return extracted


def _extract_images_from_docx(content_bytes: bytes) -> List[Dict[str, Any]]:
    extracted = []
    try:
        from docx import Document
        doc = Document(io.BytesIO(content_bytes))
        img_idx = 0
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                try:
                    img_part = rel.target_part
                    img_bytes = img_part.blob
                    if len(img_bytes) < 200:
                        continue
                    img_b64 = base64.b64encode(img_bytes).decode("utf-8")
                    ct = getattr(img_part, "content_type", "image/png")
                    extracted.append({
                        "page": 1,
                        "index": img_idx,
                        "name": f"docx_image_{img_idx}",
                        "image_b64": img_b64,
                        "mime_type": ct,
                        "size_bytes": len(img_bytes),
                        "type": "docx_image",
                    })
                    img_idx += 1
                except Exception as e:
                    print(f"[KB Ingestion] Error extracting DOCX image {img_idx}: {e}")
    except Exception as e:
        print(f"[KB Ingestion] DOCX image extraction failed: {e}")
    return extracted


def _extract_images_from_pptx(content_bytes: bytes) -> List[Dict[str, Any]]:
    extracted = []
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(io.BytesIO(content_bytes))
        for slide_num, slide in enumerate(prs.slides, 1):
            img_idx = 0
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        img_bytes = shape.image.blob
                        if len(img_bytes) < 200:
                            continue
                        img_b64 = base64.b64encode(img_bytes).decode("utf-8")
                        ct = shape.image.content_type or "image/png"
                        extracted.append({
                            "page": slide_num,
                            "index": img_idx,
                            "name": shape.name or f"slide_{slide_num}_img_{img_idx}",
                            "image_b64": img_b64,
                            "mime_type": ct,
                            "size_bytes": len(img_bytes),
                            "type": "pptx_image",
                        })
                        img_idx += 1
                    except Exception as e:
                        print(f"[KB Ingestion] Error extracting PPTX image from slide {slide_num}: {e}")
    except Exception as e:
        print(f"[KB Ingestion] PPTX image extraction failed: {e}")
    return extracted


def parse_document(file_content_b64: str, file_type: str, file_name: str) -> Tuple[str, List[Dict[str, Any]]]:
    from agents.Document_formatter_agent.tools.file_extractor import extract_file_content

    text, success, _old_images = extract_file_content(file_content_b64, file_type, file_name)
    if not success:
        raise ValueError(f"Failed to parse document: {text}")

    content_bytes = base64.b64decode(file_content_b64)
    ext = file_type.lower().strip(".")

    images_with_data = []
    if ext == "pdf":
        images_with_data = _extract_images_from_pdf(content_bytes)
    elif ext in ("docx", "doc"):
        images_with_data = _extract_images_from_docx(content_bytes)
    elif ext in ("pptx", "ppt"):
        images_with_data = _extract_images_from_pptx(content_bytes)

    print(f"[KB Ingestion] Extracted {len(images_with_data)} images with actual bytes from {file_name}")

    return text, images_with_data


def caption_document_images(images: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not images:
        return []

    images_with_data = [img for img in images if img.get("image_b64")]
    if not images_with_data:
        print("[KB Ingestion] No images with actual byte data to caption")
        return images

    api_key = os.getenv("PWC_GENAI_API_KEY") or os.getenv("GEMINI_API_KEY")
    bearer_token = os.getenv("PWC_GENAI_BEARER_TOKEN")
    raw_url = os.getenv("PWC_GENAI_ENDPOINT_URL") or os.getenv(
        "GEMINI_API_ENDPOINT",
        "https://genai-sharedservice-americas.pwc.com/completions"
    )
    chat_completions_url = raw_url.replace("/completions", "").rstrip("/") + "/chat/completions"

    if not api_key:
        print("[KB Ingestion] No GenAI API key configured, skipping image captioning")
        for img in images:
            img["caption"] = ""
        return images

    headers = {
        "accept": "application/json",
        "API-Key": api_key,
        "Content-Type": "application/json",
    }
    if bearer_token:
        headers["Authorization"] = f"Bearer {bearer_token}"

    print(f"[KB Ingestion] Using chat/completions endpoint: {chat_completions_url}")

    import concurrent.futures

    def _caption_single_image(img: Dict[str, Any]) -> Dict[str, Any]:
        try:
            mime = img.get("mime_type", "image/png")
            img_b64 = img["image_b64"]
            page_info = img.get("page", "unknown")

            request_body = {
                "model": "",
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:{mime};base64,{img_b64}"
                                }
                            },
                            {
                                "type": "text",
                                "text": (
                                    "OCR this image: extract and return ALL text visible in the image exactly as written. "
                                    "Include every word, number, label, heading, bullet point, table cell, axis label, "
                                    "legend entry, and annotation. For flowcharts and diagrams, read each box/node text "
                                    "and connection labels in logical order. For charts/graphs, extract titles, axis labels, "
                                    "data labels, and legend text. For tables, reproduce all cell content row by row. "
                                    "Do NOT summarize or describe the image — output only the raw text content found in it."
                                )
                            }
                        ]
                    }
                ],
                "temperature": 0.1,
                "max_tokens": 1000,
                "stream": False,
            }

            print(f"[KB Ingestion] Captioning image on page {page_info} ({img.get('name', 'unnamed')}, {img.get('size_bytes', 0)} bytes)...")
            resp = requests.post(chat_completions_url, json=request_body, headers=headers, timeout=60)
            resp.raise_for_status()
            data = resp.json()

            caption = ""
            if "choices" in data and data["choices"]:
                msg = data["choices"][0].get("message", {})
                caption = msg.get("content", "") if isinstance(msg, dict) else ""
            elif "content" in data:
                caption = data["content"]

            img["caption"] = caption.strip() if caption else ""
            if img["caption"]:
                print(f"[KB Ingestion] Caption OK page {page_info}: {img['caption'][:120]}...")
                try:
                    from langfuse_tracer import trace_llm_call
                    usage = data.get("usage", {})
                    trace_llm_call(
                        model="",
                        prompt=f"[image OCR page {page_info}]",
                        completion=img["caption"],
                        prompt_tokens=usage.get("prompt_tokens") or 0,
                        completion_tokens=usage.get("completion_tokens") or 0,
                        latency_ms=0,
                        agent_name="KB Ingestion Pipeline",
                        extra_metadata={"path": "vision/image_caption", "page": page_info},
                    )
                except Exception:
                    pass
            else:
                print(f"[KB Ingestion] Empty caption returned for page {page_info}")

        except Exception as e:
            print(f"[KB Ingestion] Image captioning failed for page {img.get('page', '?')}: {e}")
            img["caption"] = ""
        return img

    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as executor:
        futures = {executor.submit(_caption_single_image, img): img for img in images_with_data}
        for future in concurrent.futures.as_completed(futures):
            future.result()

    return images


def combine_text_with_captions(text: str, captioned_images: List[Dict[str, Any]]) -> str:
    if not captioned_images or not any(img.get("caption") for img in captioned_images):
        return text

    captions_by_page: Dict[int, List[str]] = {}
    for img in captioned_images:
        if img.get("caption"):
            page = img.get("page", 0)
            if page not in captions_by_page:
                captions_by_page[page] = []
            img_name = img.get("name", "image")
            captions_by_page[page].append(f"[Image OCR - {img_name}]\n{img['caption']}\n[/Image OCR]")

    if not captions_by_page:
        return text

    lines = text.split('\n')
    result_lines = []
    current_page = 1
    page_pattern = re.compile(r'--- Page (\d+) ---')

    for line in lines:
        page_match = page_pattern.match(line)
        if page_match:
            if current_page in captions_by_page:
                for caption in captions_by_page[current_page]:
                    result_lines.append(caption)
                del captions_by_page[current_page]
            current_page = int(page_match.group(1))
        result_lines.append(line)

    if current_page in captions_by_page:
        for caption in captions_by_page[current_page]:
            result_lines.append(caption)
        del captions_by_page[current_page]

    for page_num in sorted(captions_by_page.keys()):
        for caption in captions_by_page[page_num]:
            result_lines.append(caption)

    return '\n'.join(result_lines)


def chunk_text(text: str, chunk_size: int = 500) -> List[Dict[str, Any]]:
    if not text or not text.strip():
        return []

    paragraphs = re.split(r'\n\s*\n|\n(?=--- Page \d+ ---)', text)
    paragraphs = [p.strip() for p in paragraphs if p.strip()]

    chunks = []
    current_chunk = ""
    current_page = 1
    chunk_index = 0
    page_pattern = re.compile(r'--- Page (\d+) ---')

    for para in paragraphs:
        page_match = page_pattern.search(para)
        if page_match:
            current_page = int(page_match.group(1))
            para = page_pattern.sub('', para).strip()
            if not para:
                continue

        if len(current_chunk) + len(para) + 1 > chunk_size and current_chunk:
            chunks.append({
                "content": current_chunk.strip(),
                "chunkIndex": chunk_index,
                "pageNumber": current_page,
                "charCount": len(current_chunk.strip()),
            })
            chunk_index += 1
            current_chunk = ""

        if len(para) > chunk_size:
            if current_chunk:
                chunks.append({
                    "content": current_chunk.strip(),
                    "chunkIndex": chunk_index,
                    "pageNumber": current_page,
                    "charCount": len(current_chunk.strip()),
                })
                chunk_index += 1
                current_chunk = ""

            sentences = re.split(r'(?<=[.!?])\s+', para)
            for sent in sentences:
                if len(current_chunk) + len(sent) + 1 > chunk_size and current_chunk:
                    chunks.append({
                        "content": current_chunk.strip(),
                        "chunkIndex": chunk_index,
                        "pageNumber": current_page,
                        "charCount": len(current_chunk.strip()),
                    })
                    chunk_index += 1
                    current_chunk = ""
                current_chunk += (" " if current_chunk else "") + sent
        else:
            current_chunk += ("\n\n" if current_chunk else "") + para

    if current_chunk.strip():
        chunks.append({
            "content": current_chunk.strip(),
            "chunkIndex": chunk_index,
            "pageNumber": current_page,
            "charCount": len(current_chunk.strip()),
        })

    return [c for c in chunks if c["charCount"] > 30]


def extract_keywords_batch(chunks: List[Dict[str, Any]], top_n: int = 15) -> List[List[str]]:
    stop_words = {
        'the', 'a', 'an', 'and', 'or', 'but', 'in', 'on', 'at', 'to', 'for',
        'of', 'with', 'by', 'from', 'is', 'are', 'was', 'were', 'be', 'been',
        'being', 'have', 'has', 'had', 'do', 'does', 'did', 'will', 'would',
        'could', 'should', 'may', 'might', 'shall', 'can', 'need', 'must',
        'it', 'its', 'this', 'that', 'these', 'those', 'i', 'we', 'you',
        'he', 'she', 'they', 'me', 'him', 'her', 'us', 'them', 'my', 'your',
        'his', 'our', 'their', 'what', 'which', 'who', 'whom', 'when', 'where',
        'why', 'how', 'not', 'no', 'nor', 'as', 'if', 'then', 'than', 'too',
        'very', 'just', 'about', 'above', 'after', 'again', 'all', 'also',
        'any', 'because', 'before', 'between', 'both', 'each', 'few', 'more',
        'most', 'other', 'some', 'such', 'only', 'own', 'same', 'so', 'up',
        'out', 'into', 'over', 'under', 'here', 'there', 'page', 'section',
        'image', 'caption',
    }

    all_words: List[str] = []
    chunk_word_lists: List[List[str]] = []
    for chunk in chunks:
        words = re.findall(r'[a-zA-Z]{3,}', chunk["content"].lower())
        words = [w for w in words if w not in stop_words]
        chunk_word_lists.append(words)
        all_words.extend(words)

    num_chunks = len(chunks)
    doc_freq: Counter = Counter()
    for words in chunk_word_lists:
        unique = set(words)
        for w in unique:
            doc_freq[w] += 1

    result = []
    for words in chunk_word_lists:
        tf = Counter(words)
        total = len(words) if words else 1
        scored = {}
        for word, count in tf.items():
            tf_score = count / total
            idf = math.log((num_chunks + 1) / (doc_freq.get(word, 0) + 1)) + 1
            scored[word] = tf_score * idf

        top_keywords = sorted(scored, key=scored.get, reverse=True)[:top_n]
        result.append(top_keywords)

    return result


def generate_embeddings_batch(texts: List[str]) -> List[List[float]]:
    api_key = os.getenv("PWC_GENAI_API_KEY") or os.getenv("GEMINI_API_KEY")
    bearer_token = os.getenv("PWC_GENAI_BEARER_TOKEN")
    base_url = os.getenv("PWC_GENAI_ENDPOINT_URL") or os.getenv(
        "GEMINI_API_ENDPOINT",
        "https://genai-sharedservice-americas.pwc.com"
    )
    embeddings_url = base_url.replace("/completions", "").rstrip("/") + "/embeddings"

    headers = {
        "accept": "application/json",
        "API-Key": api_key,
        "Content-Type": "application/json",
    }
    if bearer_token:
        headers["Authorization"] = f"Bearer {bearer_token}"

    all_embeddings: List[List[float]] = []
    batch_size = 20

    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        payload = {
            "model": os.getenv("EMBEDDING_MODEL", "vertex_ai.text-embedding-005"),
            "input": batch,
        }
        try:
            started_at = time.time()
            resp = requests.post(embeddings_url, json=payload, headers=headers, timeout=60)
            resp.raise_for_status()
            data = resp.json()
            batch_embeddings = [item["embedding"] for item in data["data"]]

            try:
                from langfuse_tracer import trace_llm_call

                usage = data.get("usage") or {}
                total_words = sum(len((t or "").split()) for t in batch)
                avg_dim = int(sum(len(v) for v in batch_embeddings) / len(batch_embeddings)) if batch_embeddings else 0
                trace_llm_call(
                    model=str(payload.get("model") or "vertex_ai.text-embedding-005"),
                    prompt=f"[embedding batch size={len(batch)}]",
                    completion=f"[embedding batch vectors={len(batch_embeddings)} avg_dims={avg_dim}]",
                    prompt_tokens=int(usage.get("prompt_tokens") or max(1, total_words)),
                    completion_tokens=int(usage.get("completion_tokens") or 0),
                    latency_ms=(time.time() - started_at) * 1000,
                    agent_name="KB Ingestion Pipeline",
                    extra_metadata={"path": "kb_ingestion/embeddings_batch", "endpoint": "/embeddings", "batch_size": len(batch)},
                )
            except Exception:
                pass

            all_embeddings.extend(batch_embeddings)
        except Exception as e:
            print(f"[KB Ingestion] GenAI embedding batch failed: {e}, falling back to local model")
            from agents.MongoDB_RAG_agent.tools.embedding_service import get_embeddings
            local_embeddings = get_embeddings(batch)
            all_embeddings.extend(local_embeddings)

    return all_embeddings


def ingest_file_pipeline(
    mongodb_client,
    db_name: str,
    project_id: str,
    file_content_b64: str,
    file_name: str,
    file_type: str,
) -> Dict[str, Any]:

    print(f"[KB Ingestion] Step 1: File upload - {file_name}")

    print(f"[KB Ingestion] Step 2: parse_document() - extracting text + images with actual bytes")
    text, images = parse_document(file_content_b64, file_type, file_name)
    if not text.strip():
        return {"success": False, "message": "No text extracted from file", "chunks": 0}

    images_with_data = [img for img in images if img.get("image_b64")]
    print(f"[KB Ingestion] Found {len(images)} total images, {len(images_with_data)} with extractable byte data")

    print(f"[KB Ingestion] Step 3: caption_document_images() via ")
    captioned_images = caption_document_images(images)
    captioned_count = sum(1 for img in captioned_images if img.get("caption"))
    print(f"[KB Ingestion] Successfully captioned {captioned_count} images")

    print(f"[KB Ingestion] Step 4: combine_text_with_captions() - interleaving at page positions")
    combined_text = combine_text_with_captions(text, captioned_images)

    db = mongodb_client[db_name]
    docs_collection_name = f"knowledge_documents_{project_id}"
    chunks_collection_name = f"knowledge_chunks_{project_id}"
    docs_collection = db[docs_collection_name]
    chunks_collection = db[chunks_collection_name]

    print(f"[KB Ingestion] Step 5: create_knowledge_document()")
    document_id = str(uuid.uuid4())
    content_bytes = base64.b64decode(file_content_b64)

    image_summaries = []
    for img in captioned_images:
        summary = {
            "page": img.get("page"),
            "name": img.get("name", ""),
            "type": img.get("type", ""),
            "size_bytes": img.get("size_bytes", 0),
            "has_caption": bool(img.get("caption")),
            "caption_preview": (img.get("caption", "")[:200] + "...") if len(img.get("caption", "")) > 200 else img.get("caption", ""),
        }
        image_summaries.append(summary)

    doc_record = {
        "documentId": document_id,
        "filename": file_name,
        "contentType": file_type,
        "size": len(content_bytes),
        "imageCount": len(images),
        "extractedImageCount": len(images_with_data),
        "captionedImageCount": captioned_count,
        "imageSummaries": image_summaries,
        "originalFileData": file_content_b64,
        "status": "processing",
        "projectId": project_id,
        "createdAt": datetime.utcnow(),
        "updatedAt": datetime.utcnow(),
    }
    docs_collection.insert_one(doc_record)

    print(f"[KB Ingestion] Step 6: chunk_text() - ~500 char paragraph-aligned")
    chunks = chunk_text(combined_text, chunk_size=500)
    if not chunks:
        docs_collection.update_one(
            {"documentId": document_id},
            {"$set": {"status": "failed", "error": "No chunks produced", "updatedAt": datetime.utcnow()}}
        )
        return {"success": False, "message": "No chunks produced from text", "chunks": 0}

    print(f"[KB Ingestion] Step 7: extract_keywords_batch() - {len(chunks)} chunks, top 15 TF-IDF")
    keywords_batch = extract_keywords_batch(chunks, top_n=15)

    embedding_model = os.getenv("EMBEDDING_MODEL", "vertex_ai.text-embedding-005")
    print(f"[KB Ingestion] Step 8: generate_embeddings_batch() - {embedding_model} (768d)")
    chunk_texts = [c["content"] for c in chunks]
    embeddings = generate_embeddings_batch(chunk_texts)

    print(f"[KB Ingestion] Step 9: Inserting {len(chunks)} chunks into {chunks_collection_name}")
    now = datetime.utcnow()
    total_chunks = len(chunks)

    for i, (chunk, keywords, embedding) in enumerate(zip(chunks, keywords_batch, embeddings)):
        has_caption = "[Image OCR" in chunk["content"]
        chunk_doc = {
            "content": chunk["content"],
            "text": chunk["content"],
            "keywords": keywords,
            "embedding": embedding,
            "pageNumber": chunk.get("pageNumber", 1),
            "chunkIndex": chunk["chunkIndex"],
            "documentId": document_id,
            "projectId": project_id,
            "filename": file_name,
            "file_name": file_name,
            "section": f"Chunk {i + 1} of {total_chunks}",
            "charCount": chunk["charCount"],
            "hasImageCaption": has_caption,
            "keywordCount": len(keywords),
            "metadata": {
                "filename": file_name,
                "section": f"Chunk {i + 1} of {total_chunks}",
                "charCount": chunk["charCount"],
                "hasImageCaption": has_caption,
                "keywordCount": len(keywords),
                "pageNumber": chunk.get("pageNumber", 1),
            },
            "createdAt": now,
        }
        chunks_collection.insert_one(chunk_doc)

    print(f"[KB Ingestion] Step 10: update_knowledge_document() - status: ready")
    docs_collection.update_one(
        {"documentId": document_id},
        {"$set": {
            "status": "ready",
            "chunkCount": total_chunks,
            "captionedImageCount": captioned_count,
            "updatedAt": datetime.utcnow(),
        }}
    )

    print(f"[KB Ingestion] Pipeline complete: {total_chunks} chunks, {len(images)} images ({captioned_count} captioned)")

    return {
        "success": True,
        "message": f"Ingested '{file_name}' into {db_name}",
        "chunks": total_chunks,
        "file_name": file_name,
        "document_id": document_id,
        "project_id": project_id,
        "documents_collection": docs_collection_name,
        "chunks_collection": chunks_collection_name,
        "image_count": len(images),
        "extracted_images": len(images_with_data),
        "captioned_images": captioned_count,
        "embedding_dimensions": len(embeddings[0]) if embeddings else 0,
    }

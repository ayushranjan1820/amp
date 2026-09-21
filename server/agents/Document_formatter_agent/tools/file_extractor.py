import io
import base64
import re
from typing import Tuple, List, Dict, Any, Optional


def extract_from_pdf(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    extracted_images = []

    try:
        import pdfplumber
        pdf = pdfplumber.open(io.BytesIO(content))
        text_parts = []

        for page_num, page in enumerate(pdf.pages, 1):
            text = page.extract_text()
            if text:
                text_parts.append(f"--- Page {page_num} ---\n{text}")

            tables = page.extract_tables()
            for table_idx, table in enumerate(tables):
                if table:
                    table_md = _table_to_markdown(table)
                    text_parts.append(f"\n[Table {table_idx + 1} on Page {page_num}]\n{table_md}")

            if page.images:
                for img_idx, img in enumerate(page.images):
                    extracted_images.append({
                        "page": page_num,
                        "index": img_idx,
                        "bbox": [img.get("x0", 0), img.get("top", 0), img.get("x1", 0), img.get("bottom", 0)],
                        "type": "embedded_image"
                    })

        pdf.close()
        return "\n\n".join(text_parts), extracted_images

    except ImportError:
        pass

    try:
        from pypdf import PdfReader
        pdf_reader = PdfReader(io.BytesIO(content))
        text_parts = []
        for page_num, page in enumerate(pdf_reader.pages, 1):
            text = page.extract_text()
            if text:
                text_parts.append(f"--- Page {page_num} ---\n{text}")
        return "\n\n".join(text_parts), extracted_images
    except Exception as e:
        return f"Error extracting PDF content: {str(e)}", extracted_images


def _has_page_break(para) -> bool:
    from lxml import etree
    nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for br in para._element.findall('.//w:br', nsmap):
        if br.get(f'{{{nsmap["w"]}}}type') in ('page', 'column'):
            return True
    for ppr in para._element.findall('.//w:pPr/w:pageBreakBefore', nsmap):
        val = ppr.get(f'{{{nsmap["w"]}}}val')
        if val is None or val.lower() in ('true', '1', 'on'):
            return True
    return False


def _detect_section_breaks(doc) -> List[int]:
    from lxml import etree
    nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    section_para_indices = []
    for idx, para in enumerate(doc.paragraphs):
        for sect in para._element.findall('.//w:pPr/w:sectPr', nsmap):
            section_para_indices.append(idx)
    return section_para_indices


def extract_from_docx(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    extracted_images = []
    try:
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        doc = Document(io.BytesIO(content))
        text_parts = []
        current_page = 1
        section_breaks = set(_detect_section_breaks(doc))

        text_parts.append(f"--- Page {current_page} ---")

        for para_idx, para in enumerate(doc.paragraphs):
            if _has_page_break(para):
                current_page += 1
                text_parts.append(f"\n--- Page {current_page} ---")

            if para.text.strip():
                style_name = para.style.name if para.style else ""
                if "Heading" in style_name:
                    level = 1
                    for char in style_name:
                        if char.isdigit():
                            level = int(char)
                            break
                    text_parts.append(f"{'#' * level} {para.text}")
                elif "List" in style_name or "Bullet" in style_name:
                    text_parts.append(f"- {para.text}")
                else:
                    text_parts.append(para.text)

            if para_idx in section_breaks:
                current_page += 1
                text_parts.append(f"\n--- Page {current_page} ---")

        for table_idx, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                rows.append(row_data)
            if rows:
                table_md = _table_to_markdown(rows)
                text_parts.append(f"\n[Table {table_idx + 1} - Page {current_page}]\n{table_md}")

        img_idx = 0
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                extracted_images.append({
                    "index": img_idx,
                    "page": current_page,
                    "content_type": getattr(rel.target_part, 'content_type', 'unknown'),
                    "type": "embedded_image"
                })
                img_idx += 1

        return "\n\n".join(text_parts), extracted_images
    except Exception as e:
        return f"Error extracting Word document content: {str(e)}", extracted_images


def extract_from_pptx(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    extracted_images = []
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(io.BytesIO(content))
        text_parts = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"--- Slide {slide_num} ---"]

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_text.append(f"[Speaker Notes: {notes}]")

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        if para.text.strip():
                            slide_text.append(para.text.strip())

                if shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        row_data = [cell.text.strip() for cell in row.cells]
                        rows.append(row_data)
                    if rows:
                        table_md = _table_to_markdown(rows)
                        slide_text.append(f"\n[Table]\n{table_md}")

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    extracted_images.append({
                        "slide": slide_num,
                        "index": len(extracted_images),
                        "name": shape.name,
                        "type": "slide_image"
                    })

            text_parts.append("\n".join(slide_text))

        return "\n\n".join(text_parts), extracted_images
    except Exception as e:
        return f"Error extracting PowerPoint content: {str(e)}", extracted_images


def extract_from_excel(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True)
        text_parts = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames, 1):
            sheet = wb[sheet_name]
            text_parts.append(f"--- Sheet {sheet_idx}: {sheet_name} ---")

            rows = []
            for row in sheet.iter_rows():
                row_values = []
                for cell in row:
                    if cell.value is not None:
                        row_values.append(str(cell.value))
                    else:
                        row_values.append("")
                if any(v.strip() for v in row_values):
                    rows.append(row_values)

            if rows:
                table_md = _table_to_markdown(rows)
                text_parts.append(table_md)

        return "\n\n".join(text_parts), []
    except Exception as e:
        return f"Error extracting Excel content: {str(e)}", []


def extract_from_html(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    extracted_images = []
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content, 'html.parser')

        for script_or_style in soup(["script", "style"]):
            script_or_style.decompose()

        for img_idx, img in enumerate(soup.find_all('img')):
            extracted_images.append({
                "index": img_idx,
                "src": img.get('src', ''),
                "alt": img.get('alt', ''),
                "type": "html_image"
            })

        text = soup.get_text(separator='\n', strip=True)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n".join(lines), extracted_images
    except Exception as e:
        return f"Error extracting HTML content: {str(e)}", extracted_images


def extract_from_image(content: bytes, file_name: str = "") -> Tuple[str, List[Dict[str, Any]]]:
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(content))
        width, height = img.size
        img_format = img.format or "Unknown"

        b64_data = base64.b64encode(content).decode('utf-8')

        description = (
            f"[Image: {file_name or 'uploaded image'}]\n"
            f"Format: {img_format}, Size: {width}x{height} pixels, Mode: {img.mode}"
        )

        images = [{
            "index": 0,
            "format": img_format,
            "width": width,
            "height": height,
            "mode": img.mode,
            "type": "standalone_image",
            "base64_preview": b64_data[:200] + "..." if len(b64_data) > 200 else b64_data
        }]

        return description, images
    except Exception as e:
        return f"Error processing image: {str(e)}", []


def _table_to_markdown(rows: List[List]) -> str:
    if not rows:
        return ""

    max_cols = max(len(row) for row in rows)
    normalized = []
    for row in rows:
        padded = list(row) + [""] * (max_cols - len(row))
        cleaned = [str(cell).strip() if cell else "" for cell in padded]
        normalized.append(cleaned)

    col_widths = [max(len(row[i]) for row in normalized) for i in range(max_cols)]
    col_widths = [max(w, 3) for w in col_widths]

    lines = []
    header = "| " + " | ".join(normalized[0][i].ljust(col_widths[i]) for i in range(max_cols)) + " |"
    separator = "| " + " | ".join("-" * col_widths[i] for i in range(max_cols)) + " |"
    lines.append(header)
    lines.append(separator)

    for row in normalized[1:]:
        line = "| " + " | ".join(row[i].ljust(col_widths[i]) for i in range(max_cols)) + " |"
        lines.append(line)

    return "\n".join(lines)


SUPPORTED_EXTENSIONS = {
    'pdf': 'PDF Document',
    'docx': 'Word Document',
    'doc': 'Word Document (Legacy)',
    'pptx': 'PowerPoint Presentation',
    'ppt': 'PowerPoint Presentation (Legacy)',
    'xlsx': 'Excel Spreadsheet',
    'xls': 'Excel Spreadsheet (Legacy)',
    'txt': 'Plain Text',
    'text': 'Plain Text',
    'md': 'Markdown',
    'markdown': 'Markdown',
    'html': 'HTML Document',
    'htm': 'HTML Document',
    'csv': 'CSV File',
    'json': 'JSON File',
    'xml': 'XML File',
    'png': 'PNG Image',
    'jpg': 'JPEG Image',
    'jpeg': 'JPEG Image',
    'gif': 'GIF Image',
    'bmp': 'BMP Image',
    'webp': 'WebP Image',
}

IMAGE_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp'}


def extract_file_content(file_content: str, file_type: str, file_name: str = "") -> Tuple[str, bool, List[Dict[str, Any]]]:
    file_type = file_type.lower().strip('.')

    if file_type not in SUPPORTED_EXTENSIONS:
        return f"Unsupported file type: {file_type}. Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS.keys()))}", False, []

    if file_type in ['txt', 'text', 'md', 'markdown', 'csv', 'json', 'xml']:
        try:
            content = base64.b64decode(file_content).decode('utf-8')
        except Exception:
            content = file_content
        return content, True, []

    try:
        content_bytes = base64.b64decode(file_content)
    except Exception:
        return "Error: Invalid file content encoding. Expected base64.", False, []

    if file_type == 'pdf':
        text, images = extract_from_pdf(content_bytes)
    elif file_type in ['docx', 'doc']:
        text, images = extract_from_docx(content_bytes)
    elif file_type in ['pptx', 'ppt']:
        text, images = extract_from_pptx(content_bytes)
    elif file_type in ['xlsx', 'xls']:
        text, images = extract_from_excel(content_bytes)
    elif file_type in ['html', 'htm']:
        text, images = extract_from_html(content_bytes)
    elif file_type in IMAGE_EXTENSIONS:
        text, images = extract_from_image(content_bytes, file_name)
    else:
        try:
            text = content_bytes.decode('utf-8')
            images = []
        except Exception:
            return f"Unable to decode file content for type: {file_type}", False, []

    if text.startswith("Error"):
        return text, False, []

    return text, True, images

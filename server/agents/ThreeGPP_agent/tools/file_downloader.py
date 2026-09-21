import httpx
import zipfile
import io
import os
import asyncio
import tempfile
import subprocess
from typing import List, Dict, Optional
from pathlib import Path

MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 10
MAX_TEXT_LENGTH = 50000
HTTP_HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; 3GPPIndexer/1.0)"}


MAX_DOC_ENTRY_SIZE = 10 * 1024 * 1024


def _extract_doc_bytes(data: bytes, filename: str) -> str:
    if len(data) > MAX_DOC_ENTRY_SIZE:
        return ""

    ext = os.path.splitext(filename.lower())[1]

    if ext == '.docx':
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            paragraphs = []
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            if paragraphs:
                return "\n".join(paragraphs)[:MAX_TEXT_LENGTH]
        except Exception:
            pass
        return ""

    if ext == '.doc':
        try:
            with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as tmp:
                tmp.write(data)
                tmp_path = tmp.name
            try:
                result = subprocess.run(
                    ['antiword', '-w', '0', tmp_path],
                    capture_output=True, text=True, timeout=15
                )
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.strip()[:MAX_TEXT_LENGTH]
            except subprocess.TimeoutExpired:
                pass
            finally:
                os.unlink(tmp_path)
        except Exception:
            pass

        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if text:
                return text[:MAX_TEXT_LENGTH]
        except Exception:
            pass

        return ""

    if ext == '.pdf':
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages_text = []
                for i, page in enumerate(pdf.pages[:30]):
                    text = page.extract_text()
                    if text:
                        pages_text.append(f"--- Page {i+1} ---\n{text}")
                if pages_text:
                    return "\n\n".join(pages_text)[:MAX_TEXT_LENGTH]
        except Exception as e:
            return f"[Error extracting PDF: {str(e)}]"

    if ext == '.xlsx':
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
            all_text = []
            for sheet in wb.sheetnames[:5]:
                ws = wb[sheet]
                rows = []
                for row in ws.iter_rows(max_row=500, values_only=True):
                    row_text = " | ".join(str(c) for c in row if c is not None)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    all_text.append(f"=== Sheet: {sheet} ===\n" + "\n".join(rows))
            if all_text:
                return "\n\n".join(all_text)[:MAX_TEXT_LENGTH]
        except Exception:
            pass
        return ""

    if ext == '.xls':
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=data)
            all_text = []
            for sheet in wb.sheets()[:5]:
                rows = []
                for row_idx in range(min(sheet.nrows, 500)):
                    row_vals = [str(sheet.cell_value(row_idx, col)) for col in range(sheet.ncols) if sheet.cell_value(row_idx, col)]
                    row_text = " | ".join(row_vals)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    all_text.append(f"=== Sheet: {sheet.name} ===\n" + "\n".join(rows))
            if all_text:
                return "\n\n".join(all_text)[:MAX_TEXT_LENGTH]
        except Exception:
            pass
        return ""

    if ext in ('.pptx', '.ppt'):
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(data))
            slides_text = []
            for i, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
                if texts:
                    slides_text.append(f"--- Slide {i+1} ---\n" + "\n".join(texts))
            if slides_text:
                return "\n\n".join(slides_text)[:MAX_TEXT_LENGTH]
        except Exception as e:
            return f"[Error extracting PowerPoint: {str(e)}]"

    return f"[Unsupported document format: {ext}, {len(data)} bytes]"


def _extract_zip(content: bytes, url: str) -> Dict:
    texts = []
    source_files = []

    DOC_EXTENSIONS = ('.doc', '.docx', '.pdf', '.xls', '.xlsx', '.pptx', '.ppt')
    TEXT_EXTENSIONS = ('.txt', '.htm', '.html', '.csv', '.xml')

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            entries = [n for n in z.namelist() if not n.endswith('/')]
            text_entries = [n for n in entries if any(n.lower().endswith(ext) for ext in TEXT_EXTENSIONS)]
            doc_entries = [n for n in entries if any(n.lower().endswith(ext) for ext in DOC_EXTENSIONS)]

            docx_first = sorted(doc_entries, key=lambda n: (0 if n.lower().endswith('.docx') else 1 if n.lower().endswith('.pdf') else 2))
            for filename in docx_first[:MAX_ZIP_ENTRIES]:
                try:
                    data = z.read(filename)
                    extracted_text = _extract_doc_bytes(data, filename)
                    if extracted_text:
                        texts.append(f"=== {filename} ===\n{extracted_text}")
                        source_files.append(filename)
                except Exception:
                    pass

            for filename in text_entries[:MAX_ZIP_ENTRIES]:
                if len(texts) >= MAX_ZIP_ENTRIES:
                    break
                try:
                    data = z.read(filename)
                    text = data.decode('utf-8', errors='ignore')[:MAX_TEXT_LENGTH]
                    texts.append(f"=== {filename} ===\n{text}")
                    source_files.append(filename)
                except Exception:
                    pass

            if not texts:
                file_listing = "\n".join(f"  - {n}" for n in entries[:50])
                texts.append(f"ZIP contents ({len(entries)} files):\n{file_listing}")

    except Exception as e:
        texts.append(f"[Error extracting ZIP: {str(e)}]")

    return {
        "url": url,
        "text": "\n\n".join(texts),
        "source_files": source_files
    }


async def download_and_extract(url: str, extension: str) -> Optional[Dict]:
    try:
        async with httpx.AsyncClient(follow_redirects=True, headers=HTTP_HEADERS, timeout=60.0) as client:
            resp = await client.get(url)
            if resp.status_code != 200:
                return {"url": url, "text": f"[Download failed: HTTP {resp.status_code}]", "source_files": []}

            content_length = len(resp.content)
            if content_length > MAX_FILE_SIZE:
                return {
                    "url": url,
                    "text": f"[File too large: {content_length / 1024 / 1024:.1f}MB, skipped]",
                    "source_files": []
                }

            if extension == '.zip':
                return _extract_zip(resp.content, url)
            elif extension in ('.txt', '.htm', '.html', '.csv', '.xml'):
                text = resp.content.decode('utf-8', errors='ignore')[:MAX_TEXT_LENGTH]
                return {"url": url, "text": text, "source_files": [os.path.basename(url)]}
            elif extension in ('.doc', '.docx', '.pdf', '.xls', '.xlsx', '.pptx', '.ppt'):
                extracted_text = _extract_doc_bytes(resp.content, os.path.basename(url))
                return {"url": url, "text": extracted_text, "source_files": [os.path.basename(url)]}
            else:
                try:
                    text = resp.content.decode('utf-8', errors='ignore')[:MAX_TEXT_LENGTH]
                    return {"url": url, "text": text, "source_files": [os.path.basename(url)]}
                except Exception:
                    return {"url": url, "text": f"[Binary file - {content_length} bytes]", "source_files": [os.path.basename(url)]}

    except Exception as e:
        return {"url": url, "text": f"[Error downloading: {str(e)}]", "source_files": []}


async def download_multiple(files: List[Dict], max_files: int = 5) -> List[Dict]:
    tasks = []
    for f in files[:max_files]:
        tasks.append(download_and_extract(f['url'], f.get('extension', '')))
    results = await asyncio.gather(*tasks, return_exceptions=True)
    extracted = []
    for r in results:
        if isinstance(r, dict) and r.get('text'):
            extracted.append(r)
        elif isinstance(r, Exception):
            extracted.append({"url": "unknown", "text": f"[Download error: {str(r)}]", "source_files": []})
    return extracted

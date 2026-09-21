"""
File content extraction tool for BPMN Generator Agent.
Supports PDF, Word, Excel, and plain text files.
"""
import io
import base64
from typing import Optional, Tuple


def extract_from_pdf(content: bytes) -> str:
    """Extract text content from PDF file."""
    try:
        from pypdf import PdfReader
        pdf_reader = PdfReader(io.BytesIO(content))
        text_parts = []
        for page in pdf_reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n\n".join(text_parts)
    except Exception as e:
        return f"Error extracting PDF content: {str(e)}"


def extract_from_docx(content: bytes) -> str:
    """Extract text content from Word document."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return "\n\n".join(text_parts)
    except Exception as e:
        return f"Error extracting Word document content: {str(e)}"


def extract_from_excel(content: bytes) -> str:
    """Extract text content from Excel file."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True)
        text_parts = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"## Sheet: {sheet_name}")
            for row in sheet.iter_rows():
                row_values = []
                for cell in row:
                    if cell.value is not None:
                        row_values.append(str(cell.value))
                if row_values:
                    text_parts.append(" | ".join(row_values))
        return "\n".join(text_parts)
    except Exception as e:
        return f"Error extracting Excel content: {str(e)}"


def extract_file_content(file_content: str, file_type: str) -> Tuple[str, bool]:
    """
    Extract content from uploaded file based on file type.
    
    Args:
        file_content: Base64 encoded file content or plain text
        file_type: File extension (pdf, docx, xlsx, txt, etc.)
    
    Returns:
        Tuple of (extracted_text, success_flag)
    """
    file_type = file_type.lower().strip('.')
    
    if file_type in ['txt', 'text', 'md', 'markdown']:
        try:
            content = base64.b64decode(file_content).decode('utf-8')
        except:
            content = file_content
        return content, True
    
    try:
        content_bytes = base64.b64decode(file_content)
    except:
        return "Error: Invalid file content encoding", False
    
    if file_type == 'pdf':
        text = extract_from_pdf(content_bytes)
    elif file_type in ['docx', 'doc']:
        text = extract_from_docx(content_bytes)
    elif file_type in ['xlsx', 'xls']:
        text = extract_from_excel(content_bytes)
    else:
        try:
            text = content_bytes.decode('utf-8')
        except:
            return f"Unsupported file type: {file_type}", False
    
    if text.startswith("Error"):
        return text, False
    
    return text, True

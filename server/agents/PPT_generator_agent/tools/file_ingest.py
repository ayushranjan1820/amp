"""File ingestion helpers for PPT Generator Agent.

Supports extracting text from uploaded CSV, Excel, PDF, TXT, and DOCX files,
then chunking and ranking chunks for slide-level authoring.
"""

from __future__ import annotations

import base64
import csv
import io
import re
from typing import Any, Dict, List, Tuple


_TEXT_TYPES = {"txt", "text", "md", "markdown"}
_SUPPORTED_TYPES = {"csv", "xlsx", "xls", "pdf", "docx", "doc"} | _TEXT_TYPES


def normalize_file_type(file_type: str | None, file_name: str | None = None) -> str:
    ft = (file_type or "").strip().lower().lstrip(".")
    if ft:
        return ft
    if file_name and "." in file_name:
        return file_name.rsplit(".", 1)[-1].strip().lower()
    return "txt"


def _decode_text_bytes(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(enc)
        except Exception:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_from_csv(content: bytes) -> str:
    raw = _decode_text_bytes(content)
    reader = csv.reader(io.StringIO(raw))
    lines: List[str] = []
    for idx, row in enumerate(reader, start=1):
        if idx > 5000:
            lines.append("[csv truncated after 5000 rows]")
            break
        cleaned = [str(c).strip() for c in row]
        if any(cleaned):
            lines.append(f"Row {idx}: " + " | ".join(cleaned))
    return "\n".join(lines)


def _extract_from_excel(content: bytes, file_type: str) -> str:
    parts: List[str] = []

    if file_type == "xls":
        try:
            import xlrd  # type: ignore

            wb = xlrd.open_workbook(file_contents=content)
            for sheet in wb.sheets():
                parts.append(f"## Sheet: {sheet.name}")
                max_rows = min(sheet.nrows, 1200)
                for r in range(max_rows):
                    row_vals = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                    row_vals = [v for v in row_vals if v]
                    if row_vals:
                        parts.append(" | ".join(row_vals))
                if sheet.nrows > max_rows:
                    parts.append(f"[sheet truncated after {max_rows} rows]")
            return "\n".join(parts)
        except Exception:
            # Fall through to openpyxl attempt for mislabeled files.
            pass

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    for name in wb.sheetnames:
        sheet = wb[name]
        parts.append(f"## Sheet: {name}")
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            row_count += 1
            if row_count > 1200:
                parts.append("[sheet truncated after 1200 rows]")
                break
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def _text_quality_score(text: str) -> float:
    if not text:
        return 0.0
    alpha = len(re.findall(r"[A-Za-z]", text))
    digits = len(re.findall(r"\d", text))
    words = len(re.findall(r"\b\w+\b", text))
    printable = len(re.findall(r"[\x20-\x7E]", text))
    density = printable / max(1, len(text))
    return (alpha * 1.0) + (digits * 0.35) + (words * 1.25) + (density * 120.0)


def _looks_low_quality(text: str) -> bool:
    if not text:
        return True
    if len(text) < 120:
        return True
    words = re.findall(r"\b\w+\b", text)
    alpha = len(re.findall(r"[A-Za-z]", text))
    odd_symbols = len(re.findall(r"[^\w\s\.,;:!?()\[\]{}%$@#&/\\\-]", text))
    alpha_ratio = alpha / max(1, len(text))
    odd_ratio = odd_symbols / max(1, len(text))
    return len(words) < 40 or alpha_ratio < 0.22 or odd_ratio > 0.18


def _extract_from_pdf(content: bytes) -> str:
    parts_pypdf: List[str] = []
    parts_pdfplumber: List[str] = []

    # Path 1: pypdf (fast and good for many digital PDFs).
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        for i, page in enumerate(reader.pages, start=1):
            txt = (page.extract_text() or "").strip()
            if txt:
                parts_pypdf.append(f"## Page {i}\n{txt}")
    except Exception:
        parts_pypdf = []

    # Path 2: pdfplumber (often better layout recovery on complex docs).
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                txt = (page.extract_text() or "").strip()
                if txt:
                    parts_pdfplumber.append(f"## Page {i}\n{txt}")
    except Exception:
        parts_pdfplumber = []

    text_a = "\n\n".join(parts_pypdf).strip()
    text_b = "\n\n".join(parts_pdfplumber).strip()

    if text_a and text_b:
        return text_a if _text_quality_score(text_a) >= _text_quality_score(text_b) else text_b
    return text_a or text_b


def _extract_from_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts: List[str] = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    for t_idx, table in enumerate(doc.tables, start=1):
        parts.append(f"## Table {t_idx}")
        for row in table.rows:
            vals = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if vals:
                parts.append(" | ".join(vals))

    return "\n".join(parts)


def extract_uploaded_text(file_content: str, file_type: str | None, file_name: str | None = None) -> Tuple[str, Dict[str, Any]]:
    ft = normalize_file_type(file_type, file_name)
    meta: Dict[str, Any] = {
        "file_type": ft,
        "file_name": file_name or "uploaded_file",
        "supported": ft in _SUPPORTED_TYPES,
    }

    if ft not in _SUPPORTED_TYPES:
        return "", {**meta, "error": f"Unsupported file type: {ft}"}

    if ft in _TEXT_TYPES:
        try:
            text = base64.b64decode(file_content).decode("utf-8")
        except Exception:
            text = file_content
        cleaned = _clean_text(text)
        return cleaned, {
            **meta,
            "chars": len(cleaned),
            "lines": cleaned.count("\n") + 1,
            "quality_score": _text_quality_score(cleaned),
            "low_quality": _looks_low_quality(cleaned),
        }

    try:
        content = base64.b64decode(file_content)
    except Exception:
        return "", {**meta, "error": "Invalid base64 file content"}

    try:
        if ft == "csv":
            text = _extract_from_csv(content)
        elif ft in {"xlsx", "xls"}:
            text = _extract_from_excel(content, ft)
        elif ft == "pdf":
            text = _extract_from_pdf(content)
        elif ft in {"docx", "doc"}:
            text = _extract_from_docx(content)
        else:
            text = _decode_text_bytes(content)
    except Exception as e:
        return "", {**meta, "error": f"Extraction failed: {e}"}

    cleaned = _clean_text(text)
    if not cleaned:
        return "", {**meta, "error": "No extractable text found in file"}

    return cleaned, {
        **meta,
        "chars": len(cleaned),
        "lines": cleaned.count("\n") + 1,
        "quality_score": _text_quality_score(cleaned),
        "low_quality": _looks_low_quality(cleaned),
    }


def _clean_text(text: str) -> str:
    s = (text or "").replace("\x00", " ")
    s = re.sub(r"\r\n?", "\n", s)
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def chunk_text_for_parallel_processing(
    text: str,
    source_label: str,
    chunk_chars: int = 2200,
    overlap_chars: int = 240,
    max_chunks: int = 48,
) -> List[Dict[str, Any]]:
    if not text.strip():
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()]

    chunks: List[Dict[str, Any]] = []
    buf: List[str] = []
    buf_len = 0

    for para in paragraphs:
        p_len = len(para)
        projected = buf_len + (2 if buf else 0) + p_len
        if buf and projected > chunk_chars:
            chunk_text = "\n\n".join(buf).strip()
            if chunk_text:
                chunks.append(chunk_text)
            if len(chunks) >= max_chunks:
                break

            overlap = chunk_text[-overlap_chars:].strip() if overlap_chars > 0 else ""
            buf = [overlap, para] if overlap else [para]
            buf_len = sum(len(x) for x in buf) + (2 if len(buf) > 1 else 0)
        else:
            buf.append(para)
            buf_len = projected

    if len(chunks) < max_chunks and buf:
        tail = "\n\n".join(buf).strip()
        if tail:
            chunks.append(tail)

    out: List[Dict[str, Any]] = []
    for i, c in enumerate(chunks, start=1):
        out.append(
            {
                "chunk_id": i,
                "source": source_label,
                "title": f"{source_label} chunk {i}",
                "text": c,
                "chars": len(c),
            }
        )
    return out


def choose_relevant_chunks(
    chunks: List[Dict[str, Any]],
    query: str,
    slide_hint: str,
    top_k: int = 4,
) -> List[Dict[str, Any]]:
    if not chunks:
        return []

    tokens = _tokenize(f"{query} {slide_hint}")
    scored: List[Tuple[float, int, Dict[str, Any]]] = []
    max_overlap = 0

    for c in chunks:
        txt = c.get("text", "")
        low = txt.lower()
        overlap = sum(1 for t in tokens if t and t in low)
        max_overlap = max(max_overlap, overlap)
        numbers = len(re.findall(r"\b\d[\d,\.]*\b", txt))
        score = overlap * 3.0 + min(numbers, 25) * 0.10
        scored.append((score, int(c.get("chunk_id") or 0), c))

    # If there is no lexical overlap at all, avoid pseudo-random drift.
    if max_overlap == 0:
        ordered = sorted(chunks, key=lambda c: int(c.get("chunk_id") or 0))
        return ordered[: max(1, top_k)]

    scored.sort(key=lambda x: (x[0], -x[1]), reverse=True)
    top = [c for _, _, c in scored[: max(1, top_k)]]
    if not top:
        return chunks[: max(1, top_k)]
    return top


def chunks_to_research_rows(chunks: List[Dict[str, Any]], file_name: str) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    safe_name = (file_name or "uploaded_file").replace(" ", "_")
    for c in chunks:
        cid = int(c.get("chunk_id") or 0)
        txt = (c.get("text") or "").strip()
        if not txt:
            continue
        rows.append(
            {
                "title": c.get("title") or f"{file_name} chunk {cid}",
                "snippet": txt[:700],
                "content": txt,
                "url": f"uploaded://{safe_name}#chunk-{cid}",
                "date": "uploaded file",
            }
        )
    return rows


def _tokenize(text: str) -> List[str]:
    toks = re.findall(r"[a-zA-Z][a-zA-Z0-9_\-]{2,}", (text or "").lower())
    seen = set()
    out: List[str] = []
    for t in toks:
        if t in seen:
            continue
        seen.add(t)
        out.append(t)
    return out[:80]
